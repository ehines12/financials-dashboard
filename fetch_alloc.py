#!/usr/bin/env python3
"""Financials Time Allocation panel: builds web/data/alloc.json.

Sources (all free, no key): Yahoo Finance chart API (prices/volumes, Chrome UA), FINRA consolidated short
interest API, State Street NAV-history workbooks (SPDR ETF shares outstanding), and the FRED/MOVE series already
in web/data/data.json (run fetch_data.py first; it calls this script at the end).

Robustness: each ticker / SI symbol / shares-outstanding file is fetched independently; on failure the last good
copy from cache/ is used and flagged STALE. Groups whose inputs are all missing are dropped and listed.
"""
import io, json, math, os, subprocess, sys, time, datetime as dt
from concurrent.futures import ThreadPoolExecutor
from urllib.request import Request, urlopen
from urllib.parse import quote
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "web", "data", "alloc.json")
MACRO = os.path.join(ROOT, "web", "data", "data.json")
CACHE = os.path.join(ROOT, "cache")
CT = ZoneInfo("America/Chicago")
ET = ZoneInfo("America/New_York")
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/140.0.0.0 Safari/537.36")
P1 = int(dt.datetime(2015, 6, 1, tzinfo=dt.timezone.utc).timestamp())
FINRA_URL = "https://api.finra.org/data/group/otcMarket/name/consolidatedShortInterest"
SSGA = "https://www.ssga.com/us/en/intermediary/library-content/products/fund-data/etfs/us/navhist-us-en-{t}.xlsx"

# ---------------------------------------------------------------- universe
# Baskets are ILLUSTRATIVE equal-weight proxies (names commonly in the sub-$15B financials universe; market caps
# not verified from a data feed) until Eric supplies his coverage list. KBWR, LC and JHG return "symbol may be
# delisted" on Yahoo as of 2026-10-08 and are excluded.
GROUPS = [
    {"id": "KRE", "label": "Regional banks (KRE)", "tickers": ["KRE"], "si": ["KRE"], "so": "kre"},
    {"id": "IAT", "label": "Regional banks (IAT)", "tickers": ["IAT"], "si": ["IAT"]},
    {"id": "KBE", "label": "Banks, equal-wt (KBE)", "tickers": ["KBE"], "si": ["KBE"], "so": "kbe"},
    {"id": "KBWB", "label": "Large banks (KBWB)", "tickers": ["KBWB"], "si": ["KBWB"]},
    {"id": "PSCF", "label": "Small-cap financials (PSCF)", "tickers": ["PSCF"], "si": ["PSCF"]},
    {"id": "RSPF", "label": "Equal-wt S&P 500 financials (RSPF)", "tickers": ["RSPF"], "si": ["RSPF", "RYF"]},
    {"id": "KIE", "label": "Insurance, equal-wt (KIE)", "tickers": ["KIE"], "si": ["KIE"], "so": "kie"},
    {"id": "IAK", "label": "Insurance (IAK)", "tickers": ["IAK"], "si": ["IAK"]},
    {"id": "IAI", "label": "Brokers & exchanges (IAI)", "tickers": ["IAI"], "si": ["IAI"]},
    {"id": "REM", "label": "Mortgage REITs (REM)", "tickers": ["REM"], "si": ["REM"]},
    {"id": "KBWD", "label": "High-dividend financials (KBWD)", "tickers": ["KBWD"], "si": ["KBWD"]},
    {"id": "B_CONSFIN", "label": "Consumer finance basket", "basket": True,
     "tickers": ["ALLY", "OMF", "SLM", "BFH", "CACC", "ENVA", "NAVI"]},
    {"id": "B_SMINS", "label": "Small/mid insurers basket", "basket": True,
     "tickers": ["SIGI", "THG", "KMPR", "HMN", "RLI", "PLMR", "SKWD", "AGO"]},
    {"id": "B_AMBRK", "label": "Asset managers & brokers basket", "basket": True,
     "tickers": ["IVZ", "BEN", "AMG", "FHI", "VCTR", "SF", "PIPR", "EVR", "VIRT"]},
]
BENCH = [{"id": "XLF", "label": "Financials (XLF)", "si": ["XLF"], "so": "xlf"},
         {"id": "SPY", "label": "S&P 500 (SPY)"}, {"id": "IWM", "label": "Russell 2000 (IWM)"}]
STYLE = ["IWD", "IWF", "VTV", "VUG"]
PAIRS = [  # (id, long, short, label, scope)
    ("PSCF_XLF", "PSCF", "XLF", "Small-cap vs large financials (PSCF − XLF)", "fin"),
    ("KRE_KBWB", "KRE", "KBWB", "Regional vs large banks (KRE − KBWB)", "fin"),
    ("RSPF_XLF", "RSPF", "XLF", "Equal-wt vs cap-wt financials (RSPF − XLF)", "fin"),
    ("KBWD_XLF", "KBWD", "XLF", "High-dividend/value financials vs XLF (KBWD − XLF)", "fin"),
    ("XLF_SPY", "XLF", "SPY", "Financials vs market (XLF − SPY)", "mkt"),
    ("IWM_SPY", "IWM", "SPY", "Broad size: small vs large (IWM − SPY)", "mkt"),
    ("IWD_IWF", "IWD", "IWF", "Broad style: value vs growth (IWD − IWF)", "mkt"),
    ("VTV_VUG", "VTV", "VUG", "Broad style: value vs growth (VTV − VUG)", "mkt"),
]
FACTORS = [("DGS10", "10Y yield", 100.0, "per +10 bp"), ("T10Y2Y", "2s10s", 100.0, "per +10 bp"),
           ("BAMLH0A0HYM2", "HY OAS", 100.0, "per +10 bp"), ("MOVE", "MOVE", 1.0, "per +10 pts")]
PLACEHOLDERS = [
    ("Consensus estimate revisions (EPS/BVPS, 1M/3M breadth)", "Needs a data feed (e.g. FactSet, LSEG I/B/E/S, Bloomberg)."),
    ("Valuation: P/TBV, P/E (NTM), dividend yield vs history", "Needs fundamentals/estimates feed; Yahoo quoteSummary is not reliably accessible without a session."),
    ("Options skew / implied vol term structure by group", "Needs an options data feed (e.g. OptionMetrics, Cboe DataShop, ORATS)."),
    ("ETF fund flows for iShares/Invesco ETFs", "Only SPDR (State Street) publishes free daily shares-outstanding history; others need a flows feed."),
    ("Group market caps / coverage-list weighting", "Baskets are illustrative equal-weight; supply the coverage list to replace them."),
]


# ---------------------------------------------------------------- helpers
def load_cache(name):
    p = os.path.join(CACHE, name)
    try:
        with open(p) as f:
            return json.load(f)
    except Exception:
        return {}


def save_cache(name, obj):
    os.makedirs(CACHE, exist_ok=True)
    p = os.path.join(CACHE, name)
    with open(p + ".tmp", "w") as f:
        json.dump(obj, f, separators=(",", ":"))
    os.replace(p + ".tmp", p)


# circuit breaker: once a source has failed outright for several symbols in a row, skip it for the
# rest of the run (cached data is used and flagged STALE) instead of sleeping through every retry.
_FAILS = {"yahoo": 0, "finra": 0, "ssga": 0}
_TRIP = 3


_NET_ERRS = ("URLError", "TimeoutError", "timeout", "ConnectionError", "ConnectionRefusedError",
             "ConnectionResetError", "RemoteDisconnected", "OSError", "SSLError", "IncompleteRead")


def _breaker(src, ok, err=None):
    """Only network-level failures count toward tripping (a 404 / delisted ticker does not)."""
    if ok:
        _FAILS[src] = 0
    elif err and err.split(":")[0] in _NET_ERRS:
        _FAILS[src] += 1


def yahoo(sym, retries=4):
    last = None
    if _FAILS["yahoo"] >= _TRIP:
        return None, "skipped: Yahoo unreachable earlier in this run"
    for a in range(retries):
        host = "query1" if a % 2 == 0 else "query2"
        url = (f"https://{host}.finance.yahoo.com/v8/finance/chart/{quote(sym)}?period1={P1}"
               f"&period2={int(time.time())}&interval=1d&events=div%7Csplit")
        try:
            req = Request(url, headers={"User-Agent": UA, "Accept": "application/json"})
            with urlopen(req, timeout=30) as r:
                res = json.loads(r.read().decode())["chart"]["result"][0]
            if res["meta"].get("symbol") != sym:
                raise ValueError("symbol mismatch")
            tz = ZoneInfo(res["meta"].get("exchangeTimezoneName") or "America/New_York")
            q = res["indicators"]["quote"][0]
            adj = (res["indicators"].get("adjclose") or [{}])[0].get("adjclose") or q["close"]
            rows = {}
            for t, c, ac, v in zip(res["timestamp"], q["close"], adj, q["volume"]):
                if c is None or ac is None:
                    continue
                rows[dt.datetime.fromtimestamp(t, tz).date().isoformat()] = [round(c, 4), round(ac, 4), v or 0]
            now_et = dt.datetime.now(ET)
            today = now_et.date().isoformat()
            if today in rows and now_et.hour < 17:
                rows.pop(today)  # completed sessions only
            if len(rows) < 250:
                raise ValueError(f"too few rows ({len(rows)})")
            _breaker("yahoo", True)
            return rows, None
        except Exception as e:  # noqa: BLE001
            last = f"{type(e).__name__}: {e}"
            time.sleep(1.5 * (a + 1))
    _breaker("yahoo", False, last)
    return None, last


def finra(sym, retries=3):
    if _FAILS["finra"] >= _TRIP:
        return None, "skipped: FINRA unreachable earlier in this run"
    body = json.dumps({"limit": 5000, "compareFilters": [
        {"compareType": "EQUAL", "fieldName": "symbolCode", "fieldValue": sym}]}).encode()
    last = None
    for a in range(retries):
        try:
            req = Request(FINRA_URL, data=body, method="POST", headers={
                "Accept": "application/json", "Content-Type": "application/json", "User-Agent": "curl/8.5.0"})
            with urlopen(req, timeout=40) as r:
                recs = json.loads(r.read().decode() or "[]")
            if not isinstance(recs, list):
                raise ValueError(str(recs)[:200])
            out = {}
            for x in recs:
                if x.get("symbolCode") != sym or x.get("currentShortPositionQuantity") is None:
                    continue
                out[x["settlementDate"]] = [float(x["currentShortPositionQuantity"]),
                                            float(x.get("averageDailyVolumeQuantity") or 0),
                                            x.get("daysToCoverQuantity")]
            _breaker("finra", True)
            return out, None
        except Exception as e:  # noqa: BLE001
            last = f"{type(e).__name__}: {e}"
            time.sleep(2 * (a + 1))
    _breaker("finra", False, last)
    return None, last


def ssga_so(t, retries=3):
    if _FAILS["ssga"] >= _TRIP:
        return None, "skipped: State Street unreachable earlier in this run"
    import openpyxl
    last = None
    for a in range(retries):
        try:
            data = subprocess.run(["curl", "-sfL", "--max-time", "60", "-A", UA, SSGA.format(t=t)],
                                  capture_output=True, check=True).stdout
            if data[:2] != b"PK":
                raise ValueError("not xlsx")
            ws = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
            out = {}
            for r in ws.iter_rows(values_only=True):
                if not r or len(r) < 3 or r[0] is None:
                    continue
                try:
                    ds = dt.datetime.strptime(str(r[0]).strip(), "%d-%b-%Y").date().isoformat()
                except ValueError:
                    continue
                if isinstance(r[1], (int, float)) and isinstance(r[2], (int, float)):
                    out[ds] = [float(r[1]), float(r[2])]
            if len(out) < 500:
                raise ValueError("too few rows")
            _breaker("ssga", True)
            return out, None
        except Exception as e:  # noqa: BLE001
            last = f"{type(e).__name__}: {e}"
            time.sleep(2 * (a + 1))
    _breaker("ssga", False, last)
    return None, last


def fnum(x, nd=4):
    if x is None:
        return None
    try:
        x = float(x)
    except (TypeError, ValueError):
        return None
    if math.isnan(x) or math.isinf(x):
        return None
    return round(x, nd)


def asof(s, date):
    s = s.dropna()
    s = s[s.index <= date]
    return (s.index[-1], s.iloc[-1]) if len(s) else (None, None)


def minus_months(d, n):
    return (pd.Timestamp(d) - pd.DateOffset(months=n))


def period_return(tr, horizon):
    """Total return (%) from the close on/before the lookback date to the latest close."""
    tr = tr.dropna()
    if tr.empty:
        return None
    end = tr.index[-1]
    look = {"1w": end - pd.Timedelta(days=7), "1m": minus_months(end, 1), "3m": minus_months(end, 3),
            "ytd": pd.Timestamp(end.year - 1, 12, 31), "1y": minus_months(end, 12)}[horizon]
    d0, v0 = asof(tr, look)
    if d0 is None or d0 == end:
        return None
    return (tr.iloc[-1] / v0 - 1) * 100


def zscore_last(series):
    s = series.dropna()
    if len(s) < 30 or s.std() == 0:
        return None
    return (s.iloc[-1] - s.mean()) / s.std()


def pct_rank_last(series):
    s = series.dropna()
    if len(s) < 10:
        return None
    return (s <= s.iloc[-1]).mean() * 100


def weekly_last(s):
    """Last observation of each Friday-ending week, labelled with its actual date (not the week-end)."""
    s = s.dropna()
    if s.empty:
        return s
    idx = s.index.to_series().resample("W-FRI").last().dropna()
    return pd.Series(s.loc[idx.values].values, index=pd.DatetimeIndex(idx.values))


def cap(z, c=3.0):
    return None if z is None else max(-c, min(c, z))


# ---------------------------------------------------------------- main
def main():
    t0 = time.time()
    failures, stale = [], []
    all_tickers = sorted({t for g in GROUPS for t in g["tickers"]} | {b["id"] for b in BENCH} | set(STYLE))
    si_syms = sorted({s for g in GROUPS for s in g.get("si", [])} | {"XLF"}
                     | {t for g in GROUPS if g.get("basket") for t in g["tickers"]})
    so_files = sorted({g["so"] for g in GROUPS + BENCH if g.get("so")})

    price_cache, si_cache, so_cache = load_cache("alloc_prices.json"), load_cache("alloc_si.json"), load_cache("alloc_so.json")
    prices = {}
    for t in all_tickers:  # sequential + small pause: Yahoo rate-limits bursts
        rows, err = yahoo(t)
        if rows:
            prices[t] = rows; price_cache[t] = rows
        elif t in price_cache:
            prices[t] = price_cache[t]; stale.append(f"{t} prices"); failures.append(f"{t} prices: {err} (using cached copy)")
        else:
            failures.append(f"{t} prices: {err} (no cached copy; excluded)")
        time.sleep(0.35)
    with ThreadPoolExecutor(max_workers=4) as ex:
        si_res = dict(zip(si_syms, ex.map(finra, si_syms)))
        so_res = dict(zip(so_files, ex.map(ssga_so, so_files)))
    si = {}
    for s, (rows, err) in si_res.items():
        if rows:
            merged = {**si_cache.get(s, {}), **rows}  # accumulate history across runs
            si[s] = merged; si_cache[s] = merged
        elif s in si_cache:
            si[s] = si_cache[s]; stale.append(f"{s} short interest"); failures.append(f"{s} SI: {err} (cached)")
        else:
            failures.append(f"{s} SI: {err}")
    so = {}
    for s, (rows, err) in so_res.items():
        if rows:
            so[s] = rows; so_cache[s] = rows
        elif s in so_cache:
            so[s] = so_cache[s]; stale.append(f"{s.upper()} shares outstanding"); failures.append(f"{s.upper()} SO: {err} (cached)")
        else:
            failures.append(f"{s.upper()} SO: {err}")
    save_cache("alloc_prices.json", price_cache); save_cache("alloc_si.json", si_cache); save_cache("alloc_so.json", so_cache)

    # ---- frames
    def frame(t, col):
        r = prices.get(t)
        if not r:
            return None
        s = pd.Series({pd.Timestamp(k): v[col] for k, v in r.items()}).sort_index()
        return s.astype(float)

    cal = frame("SPY", 1).index if "SPY" in prices else None
    TR, CL, VOL = {}, {}, {}
    for t in prices:
        TR[t] = frame(t, 1).reindex(cal).ffill(limit=3)
        CL[t] = frame(t, 0).reindex(cal).ffill(limit=3)
        VOL[t] = frame(t, 2).reindex(cal)
    groups = []
    for g in GROUPS:
        have = [t for t in g["tickers"] if t in TR]
        if not have:
            failures.append(f"{g['id']}: no price data; group dropped"); continue
        if g.get("basket"):
            rets = pd.concat([TR[t].pct_change() for t in have], axis=1)
            ew = rets.mean(axis=1, skipna=True).fillna(0)
            first = max(TR[t].first_valid_index() for t in have)
            idx = (1 + ew[ew.index > first]).cumprod() * 100
            TR[g["id"]] = idx.reindex(cal); CL[g["id"]] = TR[g["id"]]
        groups.append({**g, "members": have, "missing_members": [t for t in g["tickers"] if t not in have]})
    gids = [g["id"] for g in groups]
    bench_ids = [b["id"] for b in BENCH if b["id"] in TR]
    xlf = TR.get("XLF")

    # ---- 1. performance
    H = ["1w", "1m", "3m", "ytd", "1y"]
    perf = []
    for gid in gids + bench_ids:
        row = {"id": gid, "label": next((x["label"] for x in groups + BENCH if x["id"] == gid), gid),
               "bench": gid in bench_ids, "asof": TR[gid].dropna().index[-1].date().isoformat(),
               "ret": {h: fnum(period_return(TR[gid], h), 2) for h in H}}
        row["rel_xlf"] = {h: (fnum(row["ret"][h] - period_return(xlf, h), 2)
                              if row["ret"][h] is not None and period_return(xlf, h) is not None else None) for h in H}
        perf.append(row)
    ranked = sorted([p for p in perf if not p["bench"] and p["rel_xlf"]["1m"] is not None],
                    key=lambda p: -p["rel_xlf"]["1m"])
    for i, p in enumerate(ranked):
        p["rank_1m"] = i + 1

    # RS vs XLF (daily, full history) and RRG (weekly)
    start5 = cal[0]  # full price history (since mid-2015) so the 10Y/Max ranges work
    series_out = {}
    for gid in gids:
        rs = (TR[gid] / xlf).dropna()
        rs = rs[rs.index >= start5]
        series_out[f"AL_RS_{gid}"] = [[d.date().isoformat(), fnum(v, 5)] for d, v in rs.items()]
    rrg = []
    for gid in gids:
        wk = weekly_last(TR[gid] / xlf)
        ratio = 100 * wk / wk.rolling(26).mean()
        mom = 100 * ratio / ratio.shift(4)
        df = pd.DataFrame({"r": ratio, "m": mom}).dropna().tail(8)
        if len(df):
            rrg.append({"id": gid, "trail": [[d.date().isoformat(), fnum(r, 3), fnum(m, 3)] for d, (r, m) in df.iterrows()]})

    # ---- 2. size & style pairs
    pairs = []
    for pid, a, b, label, scope in PAIRS:
        if a not in TR or b not in TR:
            failures.append(f"pair {pid}: missing {a if a not in TR else b}"); continue
        row = {"id": pid, "label": label, "scope": scope,
               "spread": {h: (fnum(period_return(TR[a], h) - period_return(TR[b], h), 2)
                              if period_return(TR[a], h) is not None and period_return(TR[b], h) is not None else None)
                          for h in H}}
        pairs.append(row)
        ratio = (TR[a] / TR[b]).dropna()
        ratio = ratio[ratio.index >= start5]
        series_out[f"AL_PAIR_{pid}"] = [[d.date().isoformat(), fnum(v, 5)] for d, v in ratio.items()]

    # ---- 3. positioning
    def si_frame(sym_list):
        recs = {}
        for s in sym_list:  # later symbols (e.g. old tickers) only fill gaps
            for k, v in (si.get(s) or {}).items():
                recs.setdefault(k, v)
        if not recs:
            return None
        df = pd.DataFrame.from_dict(recs, orient="index", columns=["si", "adv", "dtc"]).sort_index()
        df.index = pd.to_datetime(df.index)
        df["dtc_finra"] = df["dtc"].astype(float)  # FINRA floors days-to-cover at 1.00
        df["dtc"] = np.where(df["adv"] > 0, df["si"] / df["adv"], np.nan)  # unfloored SI / ADV
        return df

    def so_series(key):
        if key not in so:
            return None, None
        df = pd.DataFrame.from_dict(so[key], orient="index", columns=["nav", "so"]).sort_index()
        df.index = pd.to_datetime(df.index)
        return df["so"], df["nav"]

    positioning, si_hist = [], {}
    for g in groups + [dict(BENCH[0], members=["XLF"])]:
        gid = g["id"]
        if gid not in TR:
            continue
        row = {"id": gid, "label": g["label"], "basket": bool(g.get("basket"))}
        if g.get("basket"):
            frames = {t: si_frame([t]) for t in g["members"]}
            frames = {t: f for t, f in frames.items() if f is not None and len(f)}
            if not frames:
                positioning.append({**row, "na": "no FINRA data"}); continue
            all_dates = sorted(set.union(*(set(f.index) for f in frames.values())))
            need = max(3, math.ceil(0.7 * len(frames)))
            dates, notional, dtc = [], [], []
            for d in all_dates:
                have = [f for f in frames.values() if d in f.index]
                if len(have) < need:
                    continue
                dates.append(d)
                dtc.append(float(np.nanmean([f.loc[d, "dtc"] for f in have])))
                if len(have) == len(frames):  # notional only when every member reported
                    n = 0.0
                    for t, f in frames.items():
                        n += f.loc[d, "si"] * (asof(CL[t], d)[1] or np.nan)
                    notional.append(n / 1e9)
                else:
                    notional.append(np.nan)
            df = pd.DataFrame({"notional": notional, "dtc": dtc}, index=pd.DatetimeIndex(dates))
            row["si_shares"] = None
            row["members_with_si"] = len(frames)
        else:
            f = si_frame(g.get("si") or [gid])
            if f is None or f.empty:
                positioning.append({**row, "na": "no FINRA data"}); continue
            px = CL[gid]
            df = pd.DataFrame({"si": f["si"], "dtc": f["dtc"],
                               "notional": [f.loc[d, "si"] * (asof(px, d)[1] or np.nan) / 1e9 for d in f.index]},
                              index=f.index)
            so_s, nav_s = so_series(g.get("so")) if g.get("so") else (None, None)
            if so_s is not None:
                df["pct_so"] = [df.loc[d, "si"] / asof(so_s, d)[1] * 100 if asof(so_s, d)[1] else np.nan for d in df.index]
                # implied flows: sum of daily change in shares outstanding x NAV
                fl = (so_s.diff() * nav_s).dropna()
                endd = fl.index[-1]
                row["flows_1m_musd"] = fnum(fl[fl.index > minus_months(endd, 1)].sum() / 1e6, 1)
                row["flows_3m_musd"] = fnum(fl[fl.index > minus_months(endd, 3)].sum() / 1e6, 1)
                row["so_asof"] = endd.date().isoformat()
                row["so_chg_1m_pct"] = fnum((so_s.iloc[-1] / asof(so_s, minus_months(endd, 1))[1] - 1) * 100, 2)
            row["si_shares"] = fnum(df["si"].iloc[-1], 0)
            row["si_chg_prior_pct"] = fnum((df["si"].iloc[-1] / df["si"].iloc[-2] - 1) * 100, 2) if len(df) > 1 else None
        df = df.dropna(subset=["dtc"])
        last = df.iloc[-1]
        row.update({
            "settle": df.index[-1].date().isoformat(), "reports": len(df), "since": df.index[0].date().isoformat(),
            "notional_bn": fnum(last["notional"], 3), "notional_musd": fnum(last["notional"] * 1000, 2), "dtc": fnum(last["dtc"], 2),
            "notional_chg_prior_pct": fnum((df["notional"].iloc[-1] / df["notional"].iloc[-2] - 1) * 100, 2) if len(df) > 1 else None,
            "dtc_chg_prior": fnum(df["dtc"].iloc[-1] - df["dtc"].iloc[-2], 2) if len(df) > 1 else None,
            "dtc_pct": fnum(pct_rank_last(df["dtc"]), 0),
            "dtc_pct_1m_ago": fnum(pct_rank_last(df["dtc"].iloc[:-2]) if len(df) > 12 else None, 0),
            "notional_pct": fnum(pct_rank_last(df["notional"]), 0),
            "pct_so": fnum(last.get("pct_so"), 1) if "pct_so" in df else None,
            "pct_so_pct": fnum(pct_rank_last(df["pct_so"]), 0) if "pct_so" in df else None,
            "z_dtc_level": fnum(zscore_last(df["dtc"]), 2) if len(df) >= 30 else None,
            "z_dtc_chg": fnum(zscore_last(df["dtc"].diff(2)), 2) if len(df) >= 32 else None,
        })
        positioning.append(row)
        si_hist[gid] = [[d.date().isoformat(), fnum(v, 2)] for d, v in df["dtc"].items()]
        series_out[f"AL_DTC_{gid}"] = si_hist[gid]

    # ---- 4. macro sensitivity
    with open(MACRO) as f:
        md = json.load(f)["series"]
    fac = {}
    for fid, name, mult, unit in FACTORS:
        if fid in md and md[fid].get("data"):
            s = pd.Series({pd.Timestamp(k): v for k, v in md[fid]["data"]}).sort_index()
            fac[fid] = s.diff() * mult  # bp for rates/spreads, points for MOVE
        else:
            failures.append(f"factor {fid} missing from data.json")
    macro, beta_hist = [], {}
    for gid in gids:
        r = TR[gid].pct_change() * 100
        row = {"id": gid, "label": next(g["label"] for g in groups if g["id"] == gid), "f": {}}
        for fid, name, mult, unit in FACTORS:
            if fid not in fac:
                continue
            df = pd.concat([r.rename("r"), fac[fid].rename("x")], axis=1, join="inner").dropna()
            out = {}
            for w, tag in [(126, "6m"), (252, "1y")]:
                cov = df["r"].rolling(w).cov(df["x"])
                var = df["x"].rolling(w).var()
                beta = cov / var * 10
                corr = df["r"].rolling(w).corr(df["x"])
                out[f"beta_{tag}"] = fnum(beta.iloc[-1], 3)
                out[f"corr_{tag}"] = fnum(corr.iloc[-1], 2)
                if tag == "6m":
                    out["beta_6m_3m_ago"] = fnum(beta.iloc[-64], 3) if len(beta) > 64 else None
                    dchg = beta.diff(63).dropna()
                    out["beta_shift_z"] = fnum((beta.iloc[-1] - beta.iloc[-64]) / dchg.std(), 2) if len(dchg) > 100 and dchg.std() else None
                    out["beta_shift_sd"] = fnum(dchg.std(), 4) if len(dchg) > 100 else None
                    wk = weekly_last(beta)
                    wk = wk  # full history (weekly samples)
                    series_out[f"AL_BETA_{fid}_{gid}"] = [[d.date().isoformat(), fnum(v, 3)] for d, v in wk.items()]
            out["asof"] = df.index[-1].date().isoformat()
            row["f"][fid] = out
        macro.append(row)

    # ---- 5. valuation / sentiment proxies
    val = []
    for gid in gids:
        px = (CL[gid] if gid in prices else TR[gid]).dropna()
        tr = TR[gid].dropna()
        hi52 = px.rolling(252, min_periods=200).max()
        dist = (px / hi52 - 1) * 100
        gap = (px / px.rolling(200).mean() - 1) * 100
        rv = tr.pct_change().rolling(21).std() * math.sqrt(252) * 100
        g = next(x for x in groups if x["id"] == gid)
        surges = []
        for t in g["members"]:
            v = VOL[t].dropna()
            if len(v) > 70 and v.tail(63).mean() > 0:
                surges.append(v.tail(5).mean() / v.tail(63).mean())
        val.append({"id": gid, "label": g["label"], "asof": px.index[-1].date().isoformat(),
                    "dist_52w_high": fnum(dist.iloc[-1], 2), "gap_200dma": fnum(gap.iloc[-1], 2),
                    "z_gap_200dma": fnum(zscore_last(gap), 2), "rv21": fnum(rv.iloc[-1], 1),
                    "rv21_pct": fnum(pct_rank_last(rv), 0), "z_rv21": fnum(zscore_last(rv), 2),
                    "vol_surge": fnum(np.mean(surges), 2) if surges else None})

    # ---- 6. where-to-look score
    def rp_z(gid):
        rel = (TR[gid] / xlf).dropna()
        r21 = (rel / rel.shift(21) - 1) * 100
        return zscore_last(r21), r21.dropna().iloc[-1] if len(r21.dropna()) else None

    scores = []
    P = {p["id"]: p for p in positioning}
    M = {m["id"]: m for m in macro}
    V = {v["id"]: v for v in val}
    fac_name = {f[0]: f[1] for f in FACTORS}
    med_shift = {}
    for fid, *_ in FACTORS:
        ds = [m["f"][fid]["beta_6m"] - m["f"][fid]["beta_6m_3m_ago"] for m in macro
              if fid in m["f"] and m["f"][fid].get("beta_6m_3m_ago") is not None and m["f"][fid].get("beta_6m") is not None]
        if ds:
            med_shift[fid] = float(np.median(ds))
    for gid in gids:
        label = next(g["label"] for g in groups if g["id"] == gid)
        comp, why = {}, {}
        z, rel21 = rp_z(gid)
        if z is not None:
            comp["rel_perf"] = abs(cap(z))
            why["rel_perf"] = (f"1M (21d) return vs XLF {rel21:+.1f}pp (z {z:+.1f} vs own history)"
                               + (f": unusually {'strong' if z > 0 else 'weak'}" if abs(z) > 1.5 else ""))
        p = P.get(gid, {})
        zs = [abs(cap(x)) for x in (p.get("z_dtc_level"), p.get("z_dtc_chg")) if x is not None]
        if zs:
            comp["positioning"] = float(np.mean(zs))
            why["positioning"] = (f"short-interest days-to-cover {p['dtc']:.2f} at the {p['dtc_pct']:.0f}th pct since "
                                  f"{p['since'][:4]}" + (f" (was {p['dtc_pct_1m_ago']:.0f}th ~1M ago)" if p.get('dtc_pct_1m_ago') is not None else ""))
        m = M.get(gid, {}).get("f", {})
        shifts = []
        for k, v in m.items():
            if v.get("beta_6m_3m_ago") is None or not v.get("beta_shift_sd") or k not in med_shift:
                continue
            d = v["beta_6m"] - v["beta_6m_3m_ago"]
            zrel = (d - med_shift[k]) / v["beta_shift_sd"]
            shifts.append((abs(zrel), k, v, d, zrel))
        if shifts:
            zabs, k, v, d, zrel = max(shifts)
            comp["macro_beta_shift"] = min(3.0, zabs)
            unit = "per +10 pts" if k == "MOVE" else "per +10 bp"
            why["macro_beta_shift"] = (f"6M beta to {fac_name[k]} moved {v['beta_6m_3m_ago']:+.2f} → {v['beta_6m']:+.2f}% {unit} "
                                       f"over 3M vs a peer-median change of {med_shift[k]:+.2f} (relative z {zrel:+.1f})")
        vv = V.get(gid, {})
        ts = [abs(cap(x)) for x in (vv.get("z_gap_200dma"), vv.get("z_rv21")) if x is not None]
        if ts:
            comp["technical_stretch"] = float(np.mean(ts))
            why["technical_stretch"] = (f"{vv['gap_200dma']:+.1f}% vs 200-day MA (z {vv['z_gap_200dma']:+.1f}), "
                                        f"{vv['dist_52w_high']:+.1f}% from 52w high, 21d vol at {vv['rv21_pct']:.0f}th pct")
        if not comp:
            continue
        score = float(np.mean(list(comp.values())))
        order = sorted(comp, key=lambda k: -comp[k])
        reason = why[order[0]]
        if len(order) > 1 and comp[order[1]] >= 1.0:
            reason += "; also " + why[order[1]]
        scores.append({"id": gid, "label": label, "score": fnum(score, 2),
                       "components": {k: fnum(v, 2) for k, v in comp.items()}, "n_components": len(comp),
                       "top": order[0], "reason": reason})
    scores.sort(key=lambda x: -x["score"])
    for i, s in enumerate(scores):
        s["rank"] = i + 1

    now = dt.datetime.now(dt.timezone.utc)
    payload = {
        "generated_utc": now.isoformat(timespec="seconds"),
        "generated_ct": now.astimezone(CT).strftime("%Y-%m-%d %I:%M %p %Z"),
        "price_asof": cal[-1].date().isoformat(),
        "groups": [{"id": g["id"], "label": g["label"], "basket": bool(g.get("basket")), "members": g["members"],
                    "missing_members": g["missing_members"], "si_symbols": g.get("si", g["members"] if g.get("basket") else [])}
                   for g in groups],
        "benchmarks": [b for b in BENCH if b["id"] in TR], "style": [s for s in STYLE if s in TR],
        "perf": perf, "rrg": rrg, "pairs": pairs, "positioning": positioning, "macro": macro,
        "factors": [{"id": f[0], "name": f[1], "unit": f[3]} for f in FACTORS],
        "median_beta_shift": {k: fnum(v, 3) for k, v in med_shift.items()},
        "valuation": val, "scores": scores,
        "placeholders": [{"item": a, "note": b} for a, b in PLACEHOLDERS],
        "failures": failures, "stale": stale, "series": series_out,
        "excluded_tickers": {"KBWR": "Yahoo: no data, symbol may be delisted", "LC": "Yahoo: no data, symbol may be delisted",
                             "JHG": "Yahoo: no data, symbol may be delisted"},
    }
    tmp = OUT + ".tmp"
    with open(tmp, "w") as f:
        json.dump(payload, f, separators=(",", ":"), allow_nan=False)
    os.replace(tmp, OUT)
    print(f"[{payload['generated_ct']}] alloc: {len(groups)} groups, {len(prices)}/{len(all_tickers)} tickers, "
          f"{len(si)}/{len(si_syms)} SI symbols, {len(so)}/{len(so_files)} SO files, {len(failures)} issues, "
          f"{time.time() - t0:.0f}s -> {OUT}")
    for x in failures:
        print("  ", x)
    return 0


if __name__ == "__main__":
    sys.exit(main())
